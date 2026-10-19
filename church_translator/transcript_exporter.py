"""Transcript Export Service.

Supports TXT, PDF, DOCX, Markdown, and JSON exports with language alignment,
timestamps, sermon title, and status indicators.
"""
from __future__ import annotations

import html
import json
from pathlib import Path
from typing import List, Optional

from .session_manager import TranscriptSegment, SessionRecord


class TranscriptExporter:
    """Exports sermon transcripts to multiple formats cleanly."""

    @staticmethod
    def format_segments_for_export(
        segments: List[TranscriptSegment],
        languages: str = "bilingual",  # "lv", "en", "bilingual", "secondary", "all"
        include_timestamps: bool = True,
        include_status: bool = False,
    ) -> List[dict]:
        formatted = []
        for seg in segments:
            item = {
                "id": seg.segment_id,
                "timestamp": seg.timestamp_str if include_timestamps else "",
                "status": seg.status if include_status else "",
                "provider": seg.provider,
            }
            if languages == "lv":
                item["text"] = seg.latvian_text
            elif languages == "en":
                item["text"] = seg.english_translation
            elif languages == "secondary":
                item["text"] = seg.secondary_translation
            elif languages == "bilingual":
                item["lv"] = seg.latvian_text
                item["en"] = seg.english_translation
            elif languages == "all":
                item["lv"] = seg.latvian_text
                item["en"] = seg.english_translation
                item["secondary"] = seg.secondary_translation
                item["secondary_lang"] = seg.secondary_language
            formatted.append(item)
        return formatted

    @classmethod
    def export_to_txt(
        cls,
        output_path: Path,
        session: SessionRecord,
        segments: List[TranscriptSegment],
        languages: str = "bilingual",
        include_timestamps: bool = True,
        include_title: bool = True,
    ) -> None:
        lines = []
        if include_title:
            lines.append("=" * 70)
            lines.append(f"SERMON TRANSCRIPT: {session.title}")
            lines.append(f"Date & Time: {session.created_at}")
            lines.append(f"Duration: {int(session.duration_seconds // 60)}m {int(session.duration_seconds % 60)}s")
            lines.append("=" * 70)
            lines.append("")

        for seg in segments:
            ts_str = f"[{seg.timestamp_str}] " if include_timestamps and seg.timestamp_str else ""
            if languages == "lv":
                lines.append(f"{ts_str}{seg.latvian_text}\n")
            elif languages == "en":
                lines.append(f"{ts_str}{seg.english_translation}\n")
            elif languages == "secondary":
                lines.append(f"{ts_str}{seg.secondary_translation}\n")
            elif languages == "bilingual":
                lines.append(f"{ts_str}Latvian:")
                lines.append(f"  {seg.latvian_text}")
                lines.append(f"{ts_str}English:")
                lines.append(f"  {seg.english_translation}\n")
            elif languages == "all":
                lines.append(f"{ts_str}Latvian: {seg.latvian_text}")
                lines.append(f"{ts_str}English: {seg.english_translation}")
                if seg.secondary_translation:
                    lines.append(f"{ts_str}{seg.secondary_language.upper()}: {seg.secondary_translation}")
                lines.append("")

        output_path.write_text("\n".join(lines), encoding="utf-8")

    @classmethod
    def export_to_markdown(
        cls,
        output_path: Path,
        session: SessionRecord,
        segments: List[TranscriptSegment],
        languages: str = "bilingual",
        include_timestamps: bool = True,
        include_title: bool = True,
    ) -> None:
        lines = []
        if include_title:
            lines.append(f"# {session.title}\n")
            lines.append(f"**Date:** {session.created_at}  ")
            lines.append(f"**Duration:** {int(session.duration_seconds // 60)} min {int(session.duration_seconds % 60)} sec  \n")
            lines.append("---\n")

        for seg in segments:
            ts_str = f"`{seg.timestamp_str}` " if include_timestamps and seg.timestamp_str else ""
            if languages == "lv":
                lines.append(f"{ts_str}{seg.latvian_text}\n")
            elif languages == "en":
                lines.append(f"{ts_str}{seg.english_translation}\n")
            elif languages == "secondary":
                lines.append(f"{ts_str}{seg.secondary_translation}\n")
            elif languages == "bilingual":
                lines.append(f"#### {ts_str}Segment #{seg.segment_id}")
                lines.append(f"> **Latvian:** {seg.latvian_text}\n")
                lines.append(f"> **English:** {seg.english_translation}\n")
            elif languages == "all":
                lines.append(f"#### {ts_str}Segment #{seg.segment_id}")
                lines.append(f"- **Latvian:** {seg.latvian_text}")
                lines.append(f"- **English:** {seg.english_translation}")
                if seg.secondary_translation:
                    lines.append(f"- **{seg.secondary_language.upper()}:** {seg.secondary_translation}")
                lines.append("")

        output_path.write_text("\n".join(lines), encoding="utf-8")

    @classmethod
    def export_to_json(
        cls,
        output_path: Path,
        session: SessionRecord,
        segments: List[TranscriptSegment],
        languages: str = "bilingual",
        include_timestamps: bool = True,
        include_title: bool = True,
    ) -> None:
        data = {
            "session_id": session.session_id,
            "title": session.title,
            "created_at": session.created_at,
            "duration_seconds": session.duration_seconds,
            "segments": cls.format_segments_for_export(
                segments, languages, include_timestamps, include_status=True
            ),
        }
        output_path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")

    @classmethod
    def export_to_docx(
        cls,
        output_path: Path,
        session: SessionRecord,
        segments: List[TranscriptSegment],
        languages: str = "bilingual",
        include_timestamps: bool = True,
        include_title: bool = True,
    ) -> None:
        import docx
        from docx.shared import Inches, Pt, RGBColor
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = docx.Document()

        # Set standard margins
        for section in doc.sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.8)
            section.right_margin = Inches(0.8)

        if include_title:
            title_p = doc.add_heading(session.title, level=0)
            title_p.paragraph_format.space_after = Pt(4)
            sub_p = doc.add_paragraph()
            sub_run = sub_p.add_run(
                f"Date: {session.created_at} | Duration: {int(session.duration_seconds // 60)} min {int(session.duration_seconds % 60)} sec"
            )
            sub_run.font.size = Pt(10)
            sub_run.font.color.rgb = RGBColor(100, 116, 139)
            sub_p.paragraph_format.space_after = Pt(16)

        for seg in segments:
            ts_str = f"[{seg.timestamp_str}] " if include_timestamps and seg.timestamp_str else ""
            if languages == "lv":
                p = doc.add_paragraph()
                if ts_str:
                    r_ts = p.add_run(ts_str)
                    r_ts.font.color.rgb = RGBColor(100, 116, 139)
                    r_ts.font.size = Pt(9.5)
                p.add_run(seg.latvian_text)
                p.paragraph_format.space_after = Pt(6)
            elif languages == "en":
                p = doc.add_paragraph()
                if ts_str:
                    r_ts = p.add_run(ts_str)
                    r_ts.font.color.rgb = RGBColor(100, 116, 139)
                    r_ts.font.size = Pt(9.5)
                p.add_run(seg.english_translation)
                p.paragraph_format.space_after = Pt(6)
            elif languages in ("bilingual", "all"):
                # Clean bilingual card/block
                p_lv = doc.add_paragraph()
                if ts_str:
                    r_ts = p_lv.add_run(ts_str)
                    r_ts.font.color.rgb = RGBColor(100, 116, 139)
                    r_ts.font.size = Pt(9)
                r_lbl = p_lv.add_run("Latvian: ")
                r_lbl.bold = True
                r_lbl.font.color.rgb = RGBColor(13, 148, 136)
                p_lv.add_run(seg.latvian_text)
                p_lv.paragraph_format.space_after = Pt(2)

                p_en = doc.add_paragraph()
                r_lbl_en = p_en.add_run("English: ")
                r_lbl_en.bold = True
                r_lbl_en.font.color.rgb = RGBColor(5, 150, 105)
                p_en.add_run(seg.english_translation)
                p_en.paragraph_format.space_after = Pt(8)

                if languages == "all" and seg.secondary_translation:
                    p_sec = doc.add_paragraph()
                    r_lbl_sec = p_sec.add_run(f"{seg.secondary_language.upper()}: ")
                    r_lbl_sec.bold = True
                    r_lbl_sec.font.color.rgb = RGBColor(217, 119, 6)
                    p_sec.add_run(seg.secondary_translation)
                    p_sec.paragraph_format.space_after = Pt(8)

        doc.save(str(output_path))

    @classmethod
    def export_to_pdf(
        cls,
        output_path: Path,
        session: SessionRecord,
        segments: List[TranscriptSegment],
        languages: str = "bilingual",
        include_timestamps: bool = True,
        include_title: bool = True,
    ) -> None:
        from PySide6.QtGui import QTextDocument, QPdfWriter, QPageSize, QPageLayout
        from PySide6.QtCore import QMarginsF

        # Build clean HTML document
        html_parts = [
            "<html><head><style>",
            "body { font-family: 'Segoe UI', Arial, sans-serif; font-size: 11pt; color: #1E293B; line-height: 1.5; margin: 20px; }",
            "h1 { color: #0F172A; font-size: 18pt; margin-bottom: 4px; border-bottom: 2px solid #0D9488; padding-bottom: 6px; }",
            ".meta { color: #64748B; font-size: 9pt; margin-bottom: 20px; }",
            ".segment { margin-bottom: 14px; padding: 6px 10px; background-color: #F8FAFC; border-left: 3px solid #CBD5E1; }",
            ".ts { color: #94A3B8; font-size: 8.5pt; font-weight: bold; margin-bottom: 4px; }",
            ".lv-label { color: #0D9488; font-weight: bold; }",
            ".en-label { color: #059669; font-weight: bold; }",
            ".sec-label { color: #D97706; font-weight: bold; }",
            ".text-block { margin-bottom: 4px; }",
            "</style></head><body>",
        ]

        if include_title:
            html_parts.append(f"<h1>{html.escape(session.title)}</h1>")
            html_parts.append(
                f"<div class='meta'><b>Date:</b> {html.escape(session.created_at)} | "
                f"<b>Duration:</b> {int(session.duration_seconds // 60)} min {int(session.duration_seconds % 60)} sec</div>"
            )

        for seg in segments:
            ts_html = f"<div class='ts'>[{html.escape(seg.timestamp_str)}]</div>" if include_timestamps and seg.timestamp_str else ""
            if languages == "lv":
                html_parts.append(f"<div class='segment'>{ts_html}<div class='text-block'>{html.escape(seg.latvian_text)}</div></div>")
            elif languages == "en":
                html_parts.append(f"<div class='segment'>{ts_html}<div class='text-block'>{html.escape(seg.english_translation)}</div></div>")
            elif languages == "secondary":
                html_parts.append(f"<div class='segment'>{ts_html}<div class='text-block'>{html.escape(seg.secondary_translation)}</div></div>")
            elif languages == "bilingual":
                html_parts.append(
                    f"<div class='segment'>{ts_html}"
                    f"<div class='text-block'><span class='lv-label'>Latvian:</span> {html.escape(seg.latvian_text)}</div>"
                    f"<div class='text-block'><span class='en-label'>English:</span> {html.escape(seg.english_translation)}</div>"
                    f"</div>"
                )
            elif languages == "all":
                sec_line = (
                    f"<div class='text-block'><span class='sec-label'>{html.escape(seg.secondary_language.upper())}:</span> {html.escape(seg.secondary_translation)}</div>"
                    if seg.secondary_translation else ""
                )
                html_parts.append(
                    f"<div class='segment'>{ts_html}"
                    f"<div class='text-block'><span class='lv-label'>Latvian:</span> {html.escape(seg.latvian_text)}</div>"
                    f"<div class='text-block'><span class='en-label'>English:</span> {html.escape(seg.english_translation)}</div>"
                    f"{sec_line}</div>"
                )

        html_parts.append("</body></html>")
        full_html = "".join(html_parts)

        doc = QTextDocument()
        doc.setHtml(full_html)
        writer = QPdfWriter(str(output_path))
        writer.setPageSize(QPageSize(QPageSize.PageSizeId.A4))
        layout = QPageLayout(
            QPageSize(QPageSize.PageSizeId.A4),
            QPageLayout.Orientation.Portrait,
            QMarginsF(15, 15, 15, 15)
        )
        writer.setPageLayout(layout)
        doc.print_(writer)
