"""Context request delivery and completed-session GUI regressions, without APIs."""
import json
import os
from pathlib import Path
import tempfile
import time
import unittest
from dataclasses import replace
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import numpy as np
from PySide6.QtWidgets import QApplication

from church_translator.config import load_config
from church_translator.dialogs import AISummaryDialog, SessionHistoryDialog
from church_translator.engine import EngineSettings, ProcessingItem, TranslationEngine
from church_translator.glossary import Glossary
from church_translator.main import MainWindow
from church_translator.sermon_context import SermonContextManager
from church_translator.services import OpenAITranscriber, Translator
from church_translator.session_manager import SessionManager, TranscriptSegment


class ContextIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.app = QApplication.instance() or QApplication([])

    def setUp(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        self.root = Path(temp.name)
        env = patch.dict(os.environ, {'CHURCH_TRANSLATOR_APPDATA': temp.name})
        env.start()
        self.addCleanup(env.stop)
        self.cfg = replace(load_config(), openai_api_key='test', gemini_api_key='test',
                           openai_transcription_model='gpt-4o-transcribe',
                           google_application_credentials=None, openai_fallback_enabled=False)

    def context(self):
        ctx = SermonContextManager(self.root / 'context.json')
        ctx.set_text('Ķermenis un veselība\n\nVeselība ir svarīga ķermenim. '
                     'Uzturs, ēdiens un atpūta palīdz veselībai.\n\n'
                     'Lūgšana un ticība\n\nPāvils māca par ticību. 1. Korintiešiem 6:19.', save=False)
        return ctx

    def test_docx_body_order_headings_and_table_are_preserved(self):
        from docx import Document
        document = Document()
        document.add_heading('Sermon title', 0)
        document.add_paragraph('First paragraph.')
        table = document.add_table(rows=1, cols=2)
        table.cell(0, 0).text = 'Table vocabulary'
        table.cell(0, 1).text = 'Jāņa 4:12'
        document.add_heading('Second section', 1)
        document.add_paragraph('Final paragraph.')
        path = self.root / 'source.docx'
        document.save(path)
        text, title = SermonContextManager.read_docx(path)
        self.assertEqual(title, 'Sermon title')
        self.assertLess(text.index('First paragraph'), text.index('Table vocabulary'))
        self.assertLess(text.index('Jāņa'), text.index('Second section'))
        ctx = self.context()
        ctx.set_text(text, title=title)
        self.assertIn('Jāņa 4:12', ctx.bible_references)
        self.assertGreaterEqual(len(ctx.sections), 2)

    def test_replacing_document_reindexes_and_does_not_keep_old_title(self):
        ctx = self.context()
        ctx.set_text('New title\n\nAstronomy and telescopes.', save=False)
        self.assertEqual(ctx.title, 'New title')
        self.assertNotIn('veselība', ctx.recognition_guidance()['vocabulary'])
        ctx.clear()
        self.assertEqual(ctx.recognition_guidance()['vocabulary'], '')

    def test_precomputed_guidance_bootstrap_recent_section_and_topic_transition(self):
        ctx = self.context()
        first = ctx.recognition_guidance()
        self.assertTrue(first['vocabulary'])
        self.assertEqual(first['reason'], 'bootstrap_vocabulary')
        recent = ctx.recognition_guidance('Ķermeni un veselību sargājam.')
        self.assertGreaterEqual(recent['section'], 0)
        self.assertLessEqual(len(recent['vocabulary']), 600)
        unrelated = ctx.recognition_guidance('Autobuss izbrauc deviņos no pilsētas.')
        self.assertEqual(unrelated['vocabulary'], '')
        self.assertEqual(unrelated['reason'], 'recent_topic_mismatch')

    def test_related_latvian_topics_match_without_exact_words(self):
        ctx = self.context()
        hint = ctx.translation_guidance('Miesa un pārtika ir mūsu atbildība.')
        self.assertTrue(hint)
        self.assertEqual(ctx.translation_guidance('Autobuss izbrauc deviņos.'), {})

    def stt(self, model='gpt-4o-transcribe'):
        service = OpenAITranscriber(replace(self.cfg, openai_transcription_model=model), Glossary(), MagicMock())
        service._client = MagicMock()
        service._client.audio.transcriptions.create.return_value = SimpleNamespace(
            text='Ķermeni un veselību sargājam.', logprobs=[])
        return service

    def test_supported_stt_prompt_reaches_request_and_is_bounded(self):
        for model in ('gpt-4o-transcribe', 'gpt-4o-mini-transcribe', 'whisper-1'):
            with self.subTest(model=model):
                service = self.stt(model)
                service.transcribe(np.ones(16000, dtype='float32') * .05,
                                   previous_context='Ķermeni un veselību sargājam.',
                                   prepared_context=self.context().recognition_guidance())
                kwargs = service._client.audio.transcriptions.create.call_args.kwargs
                self.assertIn('reference_vocabulary', kwargs['prompt'])
                self.assertIn('previous_speech_not_current_audio', kwargs['prompt'])
                self.assertLessEqual(len(kwargs['prompt']), 1400)
                self.assertTrue(service.last_context_diagnostics['prepared_context'])

    def test_unsupported_stt_does_not_send_or_claim_prompt(self):
        service = self.stt('gpt-4o-transcribe-diarize')
        service.transcribe(np.ones(16000, dtype='float32') * .05,
                           previous_context='Earlier speech.', prepared_context=self.context().recognition_guidance())
        kwargs = service._client.audio.transcriptions.create.call_args.kwargs
        self.assertNotIn('prompt', kwargs)
        self.assertFalse(service.last_context_diagnostics['prepared_context'])
        self.assertEqual(service.last_context_diagnostics['omitted'], 'unsupported_model')

    def test_sparse_context_echo_on_music_is_checked_once_without_references(self):
        service = self.stt()
        service._client.audio.transcriptions.create.side_effect = [
            SimpleNamespace(text='Dievs rūpējas par mūsu miesu.', logprobs=[]), SimpleNamespace(text='♪ ♪', logprobs=[])]
        result = service.transcribe(np.ones(8 * 16000, dtype='float32') * .05,
                                   prepared_context=self.context().recognition_guidance())
        self.assertEqual(result.text, '')
        self.assertTrue(result.uncertain)
        calls = service._client.audio.transcriptions.create.call_args_list
        self.assertEqual(len(calls), 2)
        self.assertIn('prompt', calls[0].kwargs)
        self.assertNotIn('prompt', calls[1].kwargs)
        self.assertTrue(service.last_context_reconciliation['disagreed'])

    def test_genuine_short_spoken_phrase_is_preserved_after_context_check(self):
        service = self.stt()
        service._client.audio.transcriptions.create.return_value = SimpleNamespace(text='Jēzus Kristus.', logprobs=[])
        result = service.transcribe(np.ones(8 * 16000, dtype='float32') * .05,
                                   prepared_context=self.context().recognition_guidance())
        self.assertEqual(result.text, 'Jēzus Kristus.')
        self.assertEqual(service._client.audio.transcriptions.create.call_count, 2)

    def test_whisper_prompt_obeys_conservative_byte_budget(self):
        service = self.stt('whisper-1')
        service.transcribe(np.ones(16000, dtype='float32') * .05,
            previous_context='Ķermeni un veselību sargājam. ' * 50,
            prepared_context=self.context().recognition_guidance())
        prompt = service._client.audio.transcriptions.create.call_args.kwargs['prompt']
        self.assertLessEqual(len(prompt.encode('utf-8')), 224)
        json.loads(prompt.split('\n', 1)[1])

    def engine(self, ctx):
        with patch.object(Translator, '_init_gemini_client'):
            engine = TranslationEngine(self.cfg, EngineSettings(0, True, False, None, None, lambda: 1, lambda: 1),
                MagicMock(), MagicMock(), MagicMock(), MagicMock(), MagicMock(), sermon_context=ctx)
        for queue in (engine._chunks, engine._translations, engine._speech, engine._unrecognized):
            self.addCleanup(queue.close)
        return engine

    def test_production_engine_sends_prepared_and_only_previous_recognized_source(self):
        ctx = self.context()
        engine = self.engine(ctx)
        engine._transcriber = MagicMock()
        engine._transcriber.transcribe.return_value = SimpleNamespace(text='Ķermeni un veselību sargājam.', uncertain=False)
        for i in range(2):
            samples = np.resize(np.array([.05, -.05], dtype='float32'), 16000)
            engine._chunks.put(ProcessingItem(i, time.monotonic(), audio=samples))
        engine._chunks.finish()
        with patch('church_translator.engine.ChunkRecorder'):
            engine._process_loop()
        calls = engine._transcriber.transcribe.call_args_list
        self.assertEqual(len(calls), 2)
        self.assertIsNone(calls[0].kwargs['previous_context'])
        self.assertTrue(calls[0].kwargs['prepared_context']['vocabulary'])
        self.assertEqual(calls[1].kwargs['previous_context'], 'Ķermeni un veselību sargājam.')
        self.assertEqual(engine._translations.qsize(), 2)

    def test_low_energy_nonzero_audio_is_retained_without_billable_guess(self):
        engine = self.engine(None)
        engine._transcriber = MagicMock()
        samples = np.resize(np.array([.00001, -.00001], dtype='float32'), 3200)
        engine._chunks.put(ProcessingItem(0, time.monotonic(), audio=samples))
        engine._chunks.finish()
        with patch('church_translator.engine.ChunkRecorder'):
            engine._process_loop()
        self.assertEqual(engine._unrecognized.qsize(), 1)
        engine._transcriber.transcribe.assert_not_called()

    def test_quiet_gain_preserves_shape_and_headroom(self):
        engine = self.engine(None)
        samples = np.sin(np.arange(16000) * .17).astype('float32') * .02
        boosted, gain = engine._prepare_whisper_audio(samples, float(np.sqrt(np.mean(samples ** 2))), .02)
        self.assertGreater(gain, 2)
        self.assertLess(np.max(np.abs(boosted)), .98)
        self.assertGreater(np.corrcoef(samples, boosted)[0, 1], .99999)

    def test_completed_snapshot_and_summary_identity_survive_restart(self):
        manager = SessionManager(self.root / 'sessions')
        sid = manager.start_session('First sermon')
        manager.append_segment(TranscriptSegment(0, time.time(), '12:00', 'Actual speech.', 'Translation.'))
        manager.end_session({'capture_pending': 0, 'translation_pending': 0})
        self.assertEqual(manager.summary_session_id(), sid)
        self.assertEqual(manager.completion_snapshot(sid)['segment_count'], 1)
        reopened = SessionManager(self.root / 'sessions')
        self.assertEqual(reopened.summary_session_id(), sid)
        next_sid = reopened.start_session('Next sermon')
        self.assertEqual(reopened.summary_session_id(), next_sid)
        self.assertEqual(reopened.get_segments(next_sid), [])
        self.assertEqual(reopened.get_segments(sid)[0].english_translation, 'Translation.')

    def test_end_session_then_summary_gui_and_regeneration_and_history(self):
        window = MainWindow()
        self.addCleanup(window.close)
        window._active_config = lambda: self.cfg
        sid = window.session_manager.start_session('Controlled completed sermon')
        window.context_manager.set_text('Prepared notes must not appear in summary.', save=False)
        def drain():
            window.session_manager.append_segment(TranscriptSegment(0, time.time(), '12:00',
                'Mēs rūpējamies par veselību un ķermeni.', 'We care for our health and body.'))
        window.engine = SimpleNamespace(stop=drain, completion_info={'capture_pending': 0})
        window.stop_button.setEnabled(True)
        window.stop_button.click()
        deadline = time.monotonic() + 3
        while window._stopping and time.monotonic() < deadline:
            self.app.processEvents()
            time.sleep(.005)
        self.assertFalse(window._stopping)
        self.assertIsNone(window.engine)
        self.assertFalse(window.context_manager.is_active)
        self.assertEqual(window.session_manager.get_session(sid).state, 'completed')
        with patch('church_translator.main.AISummaryDialog') as factory:
            window.summary_button.click()
            self.assertEqual(factory.call_args.args[1].session_id, sid)
            self.assertEqual(len(factory.call_args.args[2]), 1)
        summarizer = MagicMock()
        summarizer.config = self.cfg
        summarizer.generate_summary.side_effect = ['Ķermenis un veselība.\nRūpes ir svarīgas.',
                                                   'Ķermenis un veselība.\nAtjaunots kopsavilkums.']
        summarizer.last_provider_used = 'gemini'
        summarizer.last_model_used = 'controlled-test'
        dialog = AISummaryDialog(summarizer, window.session_manager.get_session(sid),
            window.session_manager.get_segments(sid), window.session_manager, window)
        self.addCleanup(dialog.close)
        dialog.show()
        self.assertEqual(dialog.lang_combo.currentData(), 'lv')
        for i in range(2):
            dialog.generate_btn.click()
            deadline = time.monotonic() + 3
            while (dialog._is_generating or (dialog._worker_thread and dialog._worker_thread.isRunning())) and time.monotonic() < deadline:
                self.app.processEvents()
                time.sleep(.005)
            self.app.processEvents()
            self.assertFalse(dialog._is_generating)
        self.assertEqual(summarizer.generate_summary.call_count, 2)
        self.assertEqual(summarizer.generate_summary.call_args.kwargs['sermon_context'], '')
        self.assertIn('Atjaunots', window.session_manager.get_session(sid).summary_text_lv)
        history = SessionHistoryDialog(window.session_manager, window.summarizer, window)
        self.addCleanup(history.close)
        history.table.selectRow(0)
        with patch('church_translator.dialogs.AISummaryDialog') as factory:
            history._open_summary()
            self.assertEqual(factory.call_args.args[1].session_id, sid)

    def test_three_panels_outer_and_inner_geometry_equal_after_resize(self):
        window = MainWindow()
        self.addCleanup(window.close)
        window.show()
        panels = [window.latvian_text, window.english_text, window.russian_text]
        for width, height in ((1100, 760), (1600, 1000), (1280, 900)):
            window.resize(width, height)
            for _ in range(4):
                self.app.processEvents()
            for measure in (lambda p: p.parentWidget().width(), lambda p: p.parentWidget().height(),
                            lambda p: p.width(), lambda p: p.height(),
                            lambda p: p.viewport().width(), lambda p: p.viewport().height()):
                sizes = [measure(p) for p in panels]
                self.assertLessEqual(max(sizes) - min(sizes), 1, sizes)
            self.assertEqual(len({p.mapTo(window, p.rect().topLeft()).y() for p in panels}), 1)
            selector = window.secondary_language_combo
            self.assertGreaterEqual(selector.lineEdit().height(), selector.fontMetrics().height())
            headers = [p.parentWidget().layout().itemAt(0).layout() for p in panels]
            self.assertEqual(len({header.itemAt(0).widget().x() for header in headers}), 1)


if __name__ == '__main__':
    unittest.main()
